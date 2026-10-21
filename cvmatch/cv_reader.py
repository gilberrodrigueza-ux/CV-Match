from pathlib import Path


def read_cv(path):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"No encuentro el CV base: {path}")
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    elif ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in (".txt", ".md"):
        text = path.read_text(encoding="utf-8")
    else:
        raise ValueError("Formato de CV no soportado (usa .pdf, .docx o .txt)")
    text = text.strip()
    if len(text) < 300:
        raise ValueError("El CV parece vacío o escaneado como imagen; usa un PDF con texto o un .docx")
    return text
