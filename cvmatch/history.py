"""Historial de solicitudes: SQLite + exportación a Excel con hipervínculos."""
import json
import sqlite3
from datetime import date

from .config import DATA_DIR

DB_PATH = DATA_DIR / "historial.db"
ESTADOS = ["Por aplicar", "Aplicado", "Entrevista", "Prueba técnica", "Oferta", "Rechazado", "Descartado"]


def _conn():
    DATA_DIR.mkdir(exist_ok=True)
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    c.execute("""CREATE TABLE IF NOT EXISTS solicitudes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        fecha TEXT, empresa TEXT, puesto TEXT, url TEXT, ubicacion TEXT, modalidad TEXT,
        fuente TEXT, match_inicial INTEGER, match_final INTEGER, cv_docx TEXT, cv_pdf TEXT,
        reporte TEXT, jd_path TEXT, estado TEXT, notas TEXT,
        palabras_faltantes TEXT, brechas TEXT, habilidades_jd TEXT)""")
    cols = {r[1] for r in c.execute("PRAGMA table_info(solicitudes)")}
    for col in ("fecha_estado", "extras"):
        if col not in cols:
            c.execute(f"ALTER TABLE solicitudes ADD COLUMN {col} TEXT")
    return c


def add(rec):
    with _conn() as c:
        cur = c.execute("""INSERT INTO solicitudes (fecha, empresa, puesto, url, ubicacion, modalidad, fuente,
            match_inicial, match_final, cv_docx, cv_pdf, reporte, jd_path, estado, notas, palabras_faltantes,
            brechas, habilidades_jd) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
            rec.get("fecha", date.today().isoformat()), rec.get("empresa"), rec.get("puesto"), rec.get("url"),
            rec.get("ubicacion"), rec.get("modalidad"), rec.get("fuente"), rec.get("match_inicial"),
            rec.get("match_final"), rec.get("cv_docx"), rec.get("cv_pdf"), rec.get("reporte"),
            rec.get("jd_path"), rec.get("estado", "Por aplicar"), rec.get("notas", ""),
            json.dumps(rec.get("palabras_faltantes", []), ensure_ascii=False),
            json.dumps(rec.get("brechas", []), ensure_ascii=False),
            json.dumps(rec.get("habilidades_jd", []), ensure_ascii=False)))
        c.execute("UPDATE solicitudes SET fecha_estado = ?, extras = ? WHERE id = ?",
                  (rec.get("fecha", date.today().isoformat()), json.dumps(rec.get("extras", {}), ensure_ascii=False),
                   cur.lastrowid))
        return cur.lastrowid


def all_rows():
    with _conn() as c:
        return [dict(r) for r in c.execute("SELECT * FROM solicitudes ORDER BY id DESC")]


def find_by_url(url):
    if not url:
        return None
    with _conn() as c:
        r = c.execute("SELECT * FROM solicitudes WHERE url = ?", (url,)).fetchone()
        return dict(r) if r else None


def update(id_, **fields):
    if not fields:
        return
    if "estado" in fields:
        with _conn() as c:
            r = c.execute("SELECT estado FROM solicitudes WHERE id = ?", (id_,)).fetchone()
        if r and r["estado"] != fields["estado"]:
            fields["fecha_estado"] = date.today().isoformat()
    keys = ", ".join(f"{k} = ?" for k in fields)
    with _conn() as c:
        c.execute(f"UPDATE solicitudes SET {keys} WHERE id = ?", (*fields.values(), id_))


def export_excel(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = all_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "Solicitudes"
    headers = ["ID", "Fecha", "Empresa", "Puesto", "Link del trabajo", "Ubicación", "Modalidad", "Fuente",
               "Match inicial", "Match final", "CV a usar (DOCX)", "CV a usar (PDF)", "Reporte",
               "Estado", "Notas", "Palabras clave faltantes"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r["id"], r["fecha"], r["empresa"], r["puesto"], r["url"], r["ubicacion"], r["modalidad"],
                   r["fuente"], r["match_inicial"], r["match_final"], r["cv_docx"], r["cv_pdf"], r["reporte"],
                   r["estado"], r["notas"], ", ".join(json.loads(r["palabras_faltantes"] or "[]"))])
        i = ws.max_row
        for col, key in ((5, "url"), (11, "cv_docx"), (12, "cv_pdf"), (13, "reporte")):
            val = r.get(key)
            if val:
                c = ws.cell(row=i, column=col)
                c.hyperlink = val if key == "url" else "file:///" + val.replace("\\", "/")
                c.font = Font(color="0563C1", underline="single")
    widths = [5, 11, 22, 34, 40, 18, 12, 12, 9, 9, 40, 40, 40, 14, 30, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    return path


def _norm(s):
    import re
    s = (s or "").lower()
    s = re.sub(r"\b(inc|llc|ltd|s\.?a\.?|corp|corporation|costa rica|cr|the|sr|senior|ssr|jr|ii|iii)\b", " ", s)
    return " ".join(re.sub(r"[^a-z0-9áéíóúñ ]", " ", s).split())


def clave(empresa, puesto):
    return f"{_norm(empresa)}|{_norm(puesto)}"


def find_similar(empresa, puesto):
    """Solicitud previa a la misma empresa y puesto (aunque el link sea distinto)."""
    k = clave(empresa, puesto)
    for r in all_rows():
        if clave(r["empresa"], r["puesto"]) == k:
            return r
    return None


def seguimientos_pendientes(dias=7):
    """Solicitudes en 'Aplicado' sin cambio de estado hace >= dias."""
    from datetime import datetime
    out = []
    for r in all_rows():
        if r["estado"] != "Aplicado":
            continue
        f = r.get("fecha_estado") or r["fecha"]
        try:
            d = (datetime.now().date() - datetime.fromisoformat(f[:10]).date()).days
        except Exception:
            continue
        if d >= dias:
            r["dias"] = d
            out.append(r)
    return out
