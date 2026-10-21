"""Genera el CV adaptado en DOCX (ATS-friendly: una columna, sin tablas ni imágenes)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

LABELS = {
    "en": {"perfil": "Professional Profile", "competencias": "Core Competencies",
           "experiencia": "Professional Experience", "habilidades": "Technical Skills",
           "certificaciones": "Certifications", "formacion": "Training", "educacion": "Education"},
    "es": {"perfil": "Perfil Profesional", "competencias": "Competencias Clave",
           "experiencia": "Experiencia Profesional", "habilidades": "Habilidades Técnicas",
           "certificaciones": "Certificaciones", "formacion": "Formación Complementaria", "educacion": "Educación"},
}
AZUL = RGBColor(0x1F, 0x4E, 0x79)


def _lineas_atomicas(texto, ancho):
    partes = [x.strip() for x in texto.split("·") if x.strip()]
    lineas, actual = [], ""
    for parte in partes:
        cand = f"{actual} · {parte}" if actual else parte
        if len(cand) > ancho and actual:
            lineas.append(actual)
            actual = parte
        else:
            actual = cand
    if actual:
        lineas.append(actual)
    return lineas or [texto]


def _border_bottom(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F4E79")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def _para(doc, text="", size=10, bold=False, color=None, align=None, space_after=2, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = color
    return p


def _heading(doc, text):
    p = _para(doc, text.upper(), size=11, bold=True, color=AZUL, space_after=3)
    p.paragraph_format.space_before = Pt(8)
    _border_bottom(p)


def _bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def build_docx(cv, path, idioma="en"):
    L = LABELS.get(idioma, LABELS["en"])
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.55)
        s.left_margin = s.right_margin = Inches(0.65)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(10)

    _para(doc, cv.get("nombre", ""), size=20, bold=True, color=AZUL, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    if cv.get("titular"):
        _para(doc, cv["titular"], size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    if cv.get("contacto"):
        # Nunca partir un dato (link, correo) entre líneas: se agrupan por línea completos
        for linea in _lineas_atomicas(cv["contacto"], 105):
            _para(doc, linea, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        doc.paragraphs[-1].paragraph_format.space_after = Pt(4)

    if cv.get("perfil"):
        _heading(doc, L["perfil"])
        p = _para(doc, cv["perfil"], size=10, space_after=2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if cv.get("competencias"):
        _heading(doc, L["competencias"])
        _para(doc, " · ".join(cv["competencias"]), size=10)

    if cv.get("experiencia"):
        _heading(doc, L["experiencia"])
        for e in cv["experiencia"]:
            p = _para(doc, space_after=1)
            p.paragraph_format.space_before = Pt(4)
            r = p.add_run(e.get("empresa", ""))
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = AZUL
            if e.get("puesto"):
                r2 = p.add_run(f" — {e['puesto']}")
                r2.bold = True
                r2.font.size = Pt(10.5)
            meta = " | ".join(x for x in (e.get("ubicacion"), e.get("fechas")) if x)
            if meta:
                largo = len(e.get("empresa", "")) + len(e.get("puesto", "")) + len(meta) + 6
                if largo <= 100:
                    r3 = p.add_run(f" | {meta}")
                else:
                    # si no cabe, ubicación y fechas van completas en la línea siguiente (las fechas nunca se parten)
                    r3 = p.add_run()
                    r3.add_break()
                    r3.add_text(meta)
                r3.font.size = Pt(9.5)
                r3.italic = True
            for b in e.get("bullets", []):
                _bullet(doc, b)

    if cv.get("habilidades_tecnicas"):
        _heading(doc, L["habilidades"])
        for h in cv["habilidades_tecnicas"]:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(f"{h.get('categoria', '')}: ")
            r.bold = True
            r.font.size = Pt(10)
            r2 = p.add_run(h.get("items", ""))
            r2.font.size = Pt(10)

    for key, label in (("certificaciones", L["certificaciones"]), ("formacion", L["formacion"])):
        if cv.get(key):
            _heading(doc, label)
            for c in cv[key]:
                _bullet(doc, c)

    if cv.get("educacion"):
        _heading(doc, L["educacion"])
        for ed in cv["educacion"]:
            _para(doc, ed, size=10)

    doc.save(str(path))
    return path


def try_pdf(docx_path):
    """Convierte a PDF con Microsoft Word (docx2pdf). Devuelve la ruta o None."""
    try:
        from docx2pdf import convert
        pdf = str(docx_path)[:-5] + ".pdf"
        convert(str(docx_path), pdf)
        return pdf
    except Exception:
        return None


# ---------------- ajuste de páginas ----------------
LINEAS_POR_PAGINA = 52


def _l(texto, ancho=105):
    import math
    return max(1, math.ceil(len(texto or "") / ancho))


def estimate_pages(cv):
    """Estimación del número de páginas (Calibri 10, márgenes 0.65")."""
    n = 5  # nombre, titular, contacto
    if cv.get("perfil"):
        n += 2 + _l(cv["perfil"], 115)
    if cv.get("competencias"):
        n += 2 + _l(" · ".join(cv["competencias"]), 115)
    if cv.get("experiencia"):
        n += 2
        for e in cv["experiencia"]:
            n += 1.4 + sum(_l(b) + 0.1 for b in e.get("bullets", []))
    if cv.get("habilidades_tecnicas"):
        n += 2 + sum(_l(f"{h.get('categoria', '')}: {h.get('items', '')}") for h in cv["habilidades_tecnicas"])
    for k in ("certificaciones", "formacion", "educacion"):
        if cv.get(k):
            n += 2 + sum(_l(x) for x in cv[k])
    return n / LINEAS_POR_PAGINA


def fit_pages(cv, max_paginas, forzar=False):
    """Recorta bullets (sin inventar ni reescribir) empezando por el rol más antiguo hasta caber en max_paginas.
    Mínimo: 3 bullets en el rol más reciente y 1 en los demás. Con forzar=True quita 2 bullets más."""
    import copy
    cv = copy.deepcopy(cv)
    quitados = []
    exp = cv.get("experiencia") or []

    def quitar_uno():
        for i in range(len(exp) - 1, -1, -1):
            minimo = 3 if i == 0 else 1
            if len(exp[i].get("bullets", [])) > minimo:
                quitados.append(exp[i]["bullets"].pop())
                return True
        return False

    if forzar:
        for _ in range(2):
            if not quitar_uno():
                break
    else:
        while estimate_pages(cv) > max_paginas - 0.05 and quitar_uno():
            pass
    return cv, quitados


def count_pdf_pages(pdf_path):
    if not pdf_path:
        return None
    try:
        from pypdf import PdfReader
        return len(PdfReader(pdf_path).pages)
    except Exception:
        return None


def build_cover_docx(nombre, contacto, empresa, puesto, texto, path):
    from datetime import date
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.9)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    _para(doc, nombre, size=16, bold=True, color=AZUL, space_after=0)
    if contacto:
        _para(doc, contacto, size=9, space_after=10)
    _para(doc, date.today().strftime("%B %d, %Y"), size=10.5, space_after=8)
    _para(doc, f"{empresa} — {puesto}", size=10.5, bold=True, space_after=10)
    for par in [p.strip() for p in (texto or "").split("\n") if p.strip()]:
        p = _para(doc, par, size=11, space_after=8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.save(str(path))
    return path
